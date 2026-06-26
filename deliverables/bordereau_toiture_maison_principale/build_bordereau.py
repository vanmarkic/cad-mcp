#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Génère le bordereau-métré (mise hors d'eau / réparation) de la toiture de la
MAISON PRINCIPALE de la Ferme du Temple, en 3 formats à partir d'un seul modèle :
  - .csv   (séparateur ';' + BOM UTF-8  ->  ouverture directe dans Excel FR/BE)
  - .xlsx  (colonnes P.U. / Total VIDES, Total = Quantité × P.U. en formule)
  - .pdf   (métré descriptif imprimable, A4)

Les colonnes Prix Unitaire et Total sont laissées VIDES : c'est au couvreur de
les compléter pour établir son offre.

⚠️ Toutes les quantités sont PRÉSUMÉES (QP) — voir le README et le cartouche.
Régénérer :  ./.venv/bin/python deliverables/bordereau_toiture_maison_principale/build_bordereau.py
"""

import csv, os

HERE = os.path.dirname(os.path.abspath(__file__))
BASENAME = "BORDEREAU_toiture_maison_principale"

# ---------------------------------------------------------------------------
# Cartouche
# ---------------------------------------------------------------------------
META = {
    "objet": "Bordereau-métré — Mise hors d'eau / réparation de la toiture",
    "ouvrage": "Maison principale (corps de logis)",
    "chantier": "Ferme du Temple — Avenue Joseph Wauters 227, 7080 Frameries",
    "mo": "Habitat partagé de la Ferme du Temple (collectif Beaver)",
    "destinataire": "Consultation entreprises de couverture (Adritoit SRL · LB Toiture SAV)",
    "date": "13/06/2026",
    "ref": "FdT-TOIT-MP-2026-06",
    "devise": "EUR (€), montants HTVA",
}

# Données de métré (synthèse — détail/sources dans le README)
METRE = [
    ("Emprise au sol (polygone NavVis, shoelace)", "182,28 m²", "scan NavVis = méta 182,283 m²"),
    ("Surface développée des versants", "~237 m²", "estimatif architecte ; recoupt : 182/cos(39,7°) ⇒ pente ~40° (cohérent ardoise)"),
    ("Périmètre total du bâtiment", "58,40 m", "scan NavVis (shapely)"),
    ("  · arêtes LIBRES (corniche/gouttière)", "~33,6 m", "arêtes sans bâtiment mitoyen (seuil 0,8 m)"),
    ("  · arêtes MITOYENNES (solins/noues)", "~24,8 m", "contact Aile Ouest 8,8 m + Atelier 16,0 m"),
    ("Hauteur façade jusqu'à la corniche", "~10 m", "relevé MO (façade extérieure) ; cohérent scan : 8,55 − (−1,5) ≈ 10 m"),
    ("Toiture plate / terrasson zinc", "20 m²", "mesure manuelle MO sur immovision : 5 × 4 m (≪ 67 m² de l'archi)"),
    ("Couverture existante", "ardoise reconstituée", "déclaré MO ; patrimoine sans contrainte annoncée (à confirmer)"),
    ("Éléments (photos Adritoit 11/05/26)", "1 tourelle (toit conique) · 2 lucarnes jacobines (zinc cintré) · 1 œil-de-bœuf · ~3 souches · zone membrane/bâche pans N-E", "comptage photos — à confirmer"),
    ("Reprise des versants (relevé MO)", "~50 m²", "< ½ versant Nord (~68 m²) + ⅓ versant Est (~61 m²)"),
]

# ---------------------------------------------------------------------------
# Le bordereau : liste de (titre_chapitre, [ (n°, désignation, type, unité, qté_présumée) ])
#   type : QF=forfaitaire  QP=présumée  PG=poste global  SR=somme réservée  PM=pour mémoire
#   qté  : None  ->  « à relever »
# ---------------------------------------------------------------------------
CHAPTERS = [
 ("0.  INSTALLATION DE CHANTIER, ACCÈS & SÉCURITÉ", [
   ("0.1", "Installation et repli de chantier (amenée/évacuation du matériel, zone de stockage, signalisation).", "PG", "pce", 1),
   ("0.2", "Échafaudage de pied sur les façades libres concernées (montage, location, démontage), conforme à la réglementation. Base : ~33,6 m de façade libre × ~10 m.", "QP", "m² fac.", 335),
   ("0.3", "Dispositifs d'accès et de circulation en toiture (échelles de couvreur, plateforme de travail, crochets).", "QP", "forf.", 1),
   ("0.4", "Sécurité collective : garde-corps périphérique antichute, filets, et/ou ligne de vie temporaire.", "QF", "forf.", 1),
   ("0.5", "Protection des ouvrages conservés et des abords (tourelle, châssis, maçonneries, sols).", "QF", "forf.", 1),
 ]),
 ("1.  DÉPOSE & MISE HORS D'EAU PROVISOIRE", [
   ("1.1", "Dépose et évacuation des bâches provisoires existantes sur les versants.", "QP", "m²", 50),
   ("1.2", "Mise hors d'eau provisoire renforcée maintenue pendant toute la durée du chantier (bâches lestées, étanchéité temporaire).", "QF", "forf.", 1),
   ("1.3", "Sondages et diagnostic visuel des parties accessibles de la couverture et de la charpente (état des sablières / pieds de chevrons — lien avec le traitement mérule en cours).", "QF", "forf.", 1),
   ("1.4", "Dépose soignée des ardoises et de leur support dans les zones dégradées à reprendre, avec récupération des ardoises réutilisables.", "QP", "m²", 50),
   ("1.5", "Repérage / inventaire amiante préalable de la couverture existante (obligation préalable au démontage).", "QF", "forf.", 1),
   ("1.6", "Tri, descente et évacuation des déchets en CET autorisé (y c. filière amiante si avérée).", "QP", "m³", 10),
 ]),
 ("2.  CHARPENTE & SUPPORT (réparations localisées)", [
   ("2.1", "Révision générale de la charpente apparente (resserrage, calage, purge des bois pourris).", "QF", "forf.", 1),
   ("2.2", "Remplacement de chevrons dégradés par bois résineux traité classe d'emploi 2, section assortie à l'existant.", "QP", "m", 60),
   ("2.3", "Remplacement / renforcement de sablières et de pieds de chevrons (zones en contact avec maçonneries humides).", "QP", "m", 20),
   ("2.4", "Réfection du voligeage / liteaunage sous les zones reprises.", "QP", "m²", 50),
   ("2.5", "Traitement curatif et préventif du bois de charpente (fongicide / insecticide) sur parties conservées / mises à nu.", "QP", "m²", 50),
 ]),
 ("3.  SOUS-TOITURE & COUVERTURE (reprises)", [
   ("3.1", "Fourniture et pose d'un écran de sous-toiture HPV (pare-pluie) sous les zones reprises.", "QP", "m²", 50),
   ("3.2", "Contre-lattage + lattage neufs (résineux traité) sous les zones reprises.", "QP", "m²", 50),
   ("3.3", "Recouverture en ardoise RECONSTITUÉE (recomposée) assortie à l'existant sur les zones reprises, pose et fixation selon les règles de l'art. (Variante ardoise naturelle : voir poste V1.)", "QP", "m²", 50),
   ("3.4", "Remplacement ponctuel d'ardoises isolées cassées / glissées hors zones reprises.", "QP", "pce", 200),
   ("3.5", "Rétablissement du faîtage et des arêtiers (closoirs ventilés, fixation mécanique).", "QP", "m", 35),
   ("3.6", "Réfection des noues (façonnage zinc ou plomb).", "QP", "m", 12),
 ]),
 ("4.  POINTS SINGULIERS", [
   ("4.1", "Lucarnes jacobines (toit cintré zinc) : réfection couverture + joues + zinguerie (par unité).", "QP", "pce", 2),
   ("4.2", "Œil-de-bœuf / lucarne ronde : réfection de la couverture et des raccords.", "QP", "pce", 1),
   ("4.3", "Tourelle : révision / réfection de la couverture conique (plomb, zinc ou ardoise) et de son amortissement.", "PG", "forf.", 1),
   ("4.4", "Souches de cheminée : réfection des solins et bavettes en plomb, rejointoiement de la tête (par souche).", "QP", "pce", 3),
   ("4.5", "Solins et raccords d'étanchéité contre maçonneries mitoyennes (Aile Ouest, Atelier). Base : ~24,8 m d'arêtes mitoyennes.", "QP", "m", 25),
   ("4.6", "Reprise des pénétrations diverses (ventilations, sorties, scellements).", "QP", "pce", 5),
 ]),
 ("5.  ÉVACUATION DES EAUX PLUVIALES (zinguerie)", [
   ("5.1", "Dépose des gouttières / corniches et descentes dégradées (linéaire de corniche libre).", "QP", "m", 34),
   ("5.2", "Fourniture et pose de gouttières (corniche) en zinc, y c. naissances. Base : ~33,6 m d'arêtes libres.", "QP", "m", 34),
   ("5.3", "Descentes d'eaux pluviales + accessoires (dauphins fonte, crapaudines, raccords). Base : ~3 descentes × ~8 m.", "QP", "m", 24),
   ("5.4", "Bavettes, noquets et solins en zinc / plomb.", "QP", "m", 30),
 ]),
 ("6.  TOITURE PLATE / TERRASSON ZINC (sommet du comble — photos Adritoit)", [
   ("6.1", "Dépose de la couverture métallique existante (zinc / plomb) du terrasson, fortement dégradée.", "QP", "m²", 20),
   ("6.2", "Révision / remplacement du support (voligeage, platelage) de la toiture plate.", "QP", "m²", 20),
   ("6.3", "Nouvelle étanchéité ZINC à joints debout (reprise du système d'origine — solution de base). Variantes bitume/EPDM : postes V2/V3.", "QP", "m²", 20),
   ("6.4", "Relevés d'étanchéité périphériques, bavettes et raccords sur acrotères / souches (pourtour ~5×4 m).", "QP", "m", 18),
   ("6.5", "Évacuations d'eaux pluviales du terrasson (naissances, trop-pleins, raccord aux descentes).", "QP", "pce", 2),
 ]),
 ("7.  DIVERS & FIN DE CHANTIER", [
   ("7.1", "Coordination sécurité-santé / plan particulier de sécurité (PSS) — si requis.", "QF", "forf.", 1),
   ("7.2", "Nettoyage du chantier et évacuation finale des déchets.", "QF", "forf.", 1),
   ("7.3", "Réception, dossier de fin de chantier (photos, fiches techniques, garanties).", "QF", "forf.", 1),
   ("7.4", "Somme réservée pour imprévus (travaux non décelables avant ouverture) — à valider par le maître d'ouvrage.", "SR", "forf.", 1),
 ]),
]

# Variantes & options — À COMPARER, NON additionnées au total ci-dessus
# (le maître d'ouvrage retient une solution en fonction des prix).
VARIANTES = [
   ("V1", "Plus-value ardoise NATURELLE au lieu de reconstituée, sur les zones reprises (comparaison de coût demandée par le MO). À chiffrer en alternative au poste 3.3.", "QP", "m²", 50),
   ("V2", "Toiture plate — étanchéité BITUMINEUSE SBS bicouche soudée, en lieu et place du zinc (alternative au poste 6.3).", "QP", "m²", 20),
   ("V3", "Toiture plate — étanchéité EPDM monocouche collée, en lieu et place du zinc (alternative au poste 6.3).", "QP", "m²", 20),
   ("O1", "OPTION — Réfection COMPLÈTE des versants en ardoise reconstituée (dépose totale + sous-toiture + lattage + ardoises neuves), en remplacement des reprises ponctuelles. Scénario 'fin de vie' pour le financement.", "QP", "m²", 237),
   ("O2", "OPTION — idem en ardoise NATURELLE (comparaison de coût).", "QP", "m²", 237),
]

COLS = ["N°", "Désignation des ouvrages", "Type", "Unité", "Q.P.", "P.U. HTVA (€)", "Total HTVA (€)"]

def qstr(q):
    if q is None:
        return "à relever"
    return str(q).replace(".", ",")

# ===========================================================================
# 1) CSV  (Excel FR/BE : séparateur ';', BOM UTF-8)
# ===========================================================================
def build_csv(path):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow([META["objet"]])
        w.writerow([f'{META["ouvrage"]} — {META["chantier"]}'])
        w.writerow([f'Réf. {META["ref"]}  |  {META["date"]}  |  Destinataire : {META["destinataire"]}  |  {META["devise"]}'])
        w.writerow(["QUANTITÉS PRÉSUMÉES (QP) — à vérifier contradictoirement sur site. Décompte final sur métré réel."])
        w.writerow(["Surface de reprise ≈ 50 m² — relevé MO : < 1/2 du versant Nord (~68 m²) + 1/3 du versant Est (~61 m²) ; versants répartis depuis 237 m² selon l'orientation des arêtes (scan). À confirmer sur site."])
        w.writerow(["Type : QF=forfaitaire  QP=présumée  PG=poste global  SR=somme réservée  PM=pour mémoire"])
        w.writerow(["Variantes V1-V3 et options O1-O2 : ALTERNATIVES à comparer, NON additionnées au total."])
        w.writerow([])
        w.writerow(COLS)
        for title, rows in CHAPTERS:
            w.writerow([title])
            for n, d, t, u, q in rows:
                w.writerow([n, d, t, u, qstr(q), "", ""])
        w.writerow([])
        w.writerow(["", "TOTAL HTVA (chapitres 0 à 7)", "", "", "", "", ""])
        w.writerow(["", "TVA (6% rénovation logement >10 ans, sous conditions / sinon 21%)", "", "", "", "", ""])
        w.writerow(["", "TOTAL TVAC", "", "", "", "", ""])
        w.writerow([])
        w.writerow(["VARIANTES & OPTIONS — à comparer, NON incluses dans le total ci-dessus :"])
        for n, d, t, u, q in VARIANTES:
            w.writerow([n, d, t, u, qstr(q), "", ""])
    return path

# ===========================================================================
# 2) XLSX  (P.U. & Total vides ; Total = Q.P. × P.U. en formule)
# ===========================================================================
def build_xlsx(path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook(); ws = wb.active; ws.title = "Bordereau toiture"
    thin = Side(style="thin", color="BBBBBB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="1F3864")
    chap_fill = PatternFill("solid", fgColor="D9E1F2")
    tot_fill  = PatternFill("solid", fgColor="FCE4D6")
    wrap = Alignment(wrap_text=True, vertical="top")
    center = Alignment(horizontal="center", vertical="top")
    right = Alignment(horizontal="right", vertical="top")

    widths = [6, 64, 7, 9, 9, 14, 16]
    for i, wdt in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = wdt

    r = 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    c = ws.cell(r, 1, META["objet"]); c.font = Font(bold=True, size=13); r += 1
    for line in [f'{META["ouvrage"]} — {META["chantier"]}',
                 f'Maître d\'ouvrage : {META["mo"]}',
                 f'Destinataire : {META["destinataire"]}   |   Réf. {META["ref"]}   |   Date : {META["date"]}   |   {META["devise"]}']:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
        ws.cell(r, 1, line).font = Font(size=10); r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    c = ws.cell(r, 1, "⚠ QUANTITÉS PRÉSUMÉES (QP) — à vérifier contradictoirement lors d'une visite sur site. "
                      "Le décompte final est établi sur métré réel. Colonnes P.U. et Total à compléter par le couvreur. "
                      "Surface de reprise ≈ 50 m² — relevé MO : < 1/2 du versant Nord (~68 m²) + 1/3 du versant Est (~61 m²) ; "
                      "versants répartis depuis 237 m² selon l'orientation des arêtes (scan). À confirmer sur site.")
    c.font = Font(italic=True, color="C00000", size=9); c.alignment = wrap; ws.row_dimensions[r].height = 48; r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    ws.cell(r, 1, "Type : QF=forfaitaire · QP=présumée · PG=poste global · SR=somme réservée · PM=pour mémoire").font = Font(italic=True, size=8); r += 1
    r += 1

    # en-tête tableau
    hdr = r
    for i, name in enumerate(COLS, 1):
        c = ws.cell(r, i, name); c.font = Font(bold=True, color="FFFFFF"); c.fill = head_fill
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center"); c.border = border
    r += 1

    first_data = r
    for title, rows in CHAPTERS:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
        c = ws.cell(r, 1, title); c.font = Font(bold=True); c.fill = chap_fill; c.border = border; r += 1
        for n, d, t, u, q in rows:
            ws.cell(r, 1, n).alignment = center
            ws.cell(r, 2, d).alignment = wrap
            ws.cell(r, 3, t).alignment = center
            ws.cell(r, 4, u).alignment = center
            ws.cell(r, 5, (q if q is not None else "à relever")).alignment = right
            ws.cell(r, 6, None).alignment = right                    # P.U. — à compléter
            # Total = Q.P. × P.U. (si quantité numérique)
            if q is not None:
                ws.cell(r, 7, f"=IF($F{r}=\"\",\"\",E{r}*F{r})").alignment = right
            for col in range(1, 8):
                cell = ws.cell(r, col); cell.border = border
                if col == 2 and ws.row_dimensions[r].height is None:
                    ws.row_dimensions[r].height = 26
            r += 1
    last_data = r - 1

    # récap
    r += 1
    def recap(label, formula, fill=True):
        nonlocal r
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        c = ws.cell(r, 1, label); c.font = Font(bold=True); c.alignment = right
        if fill: c.fill = tot_fill
        cc = ws.cell(r, 7, formula); cc.font = Font(bold=True); cc.alignment = right
        if fill: cc.fill = tot_fill
        cc.number_format = u'#,##0.00 €'
        for col in range(1, 8): ws.cell(r, col).border = border
        r += 1
    recap("TOTAL HTVA (chapitres 0 à 7)", f"=SUM(G{first_data}:G{last_data})")
    recap("TVA 6% (rénovation logement >10 ans, sous conditions)", f"=G{last_data+2}*0.06")
    tva_row = r - 1
    recap("TOTAL TVAC", f"=G{last_data+2}+G{tva_row}")

    # --- Variantes & options : NON additionnées au total ci-dessus ---
    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    c = ws.cell(r, 1, "VARIANTES & OPTIONS — à comparer, NON incluses dans le TOTAL ci-dessus "
                      "(le maître d'ouvrage retient une solution selon les prix)")
    c.font = Font(bold=True, color="FFFFFF"); c.fill = head_fill; c.alignment = wrap
    for col in range(1, 8): ws.cell(r, col).border = border
    r += 1
    for n, d, t, u, q in VARIANTES:
        ws.cell(r, 1, n).alignment = center
        ws.cell(r, 2, d).alignment = wrap
        ws.cell(r, 3, t).alignment = center
        ws.cell(r, 4, u).alignment = center
        ws.cell(r, 5, q).alignment = right
        ws.cell(r, 6, None).alignment = right
        ws.cell(r, 7, f"=IF($F{r}=\"\",\"\",E{r}*F{r})").alignment = right   # total ligne (info, hors somme)
        for col in range(1, 8):
            ws.cell(r, col).border = border
            if col == 2 and ws.row_dimensions[r].height is None:
                ws.row_dimensions[r].height = 30
        r += 1

    ws.freeze_panes = ws.cell(first_data, 1)
    ws.print_title_rows = f"{hdr}:{hdr}"
    ws.page_setup.orientation = "portrait"; ws.page_setup.fitToWidth = 1; ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(path)
    return path

# ===========================================================================
# 3) PDF  (reportlab Platypus — A4, imprimable)
# ===========================================================================
def build_pdf(path):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import (SimpleDocTemplate, Table, TableStyle, Paragraph,
                                    Spacer, KeepTogether)

    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["Normal"], fontSize=8, leading=9.5)
    small = ParagraphStyle("small", parent=styles["Normal"], fontSize=7, leading=8.5, textColor=colors.HexColor("#555555"))
    h1 = ParagraphStyle("h1", parent=styles["Normal"], fontSize=13, leading=15, spaceAfter=2)
    warn = ParagraphStyle("warn", parent=styles["Normal"], fontSize=8, leading=10, textColor=colors.HexColor("#C00000"))
    chap = ParagraphStyle("chap", parent=styles["Normal"], fontSize=8.5, leading=10, textColor=colors.white)

    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=14*mm, rightMargin=12*mm,
                            topMargin=12*mm, bottomMargin=12*mm,
                            title=META["objet"], author=META["mo"])
    story = []
    story.append(Paragraph(META["objet"], h1))
    story.append(Paragraph(f'<b>{META["ouvrage"]}</b> — {META["chantier"]}', body))
    story.append(Paragraph(f'Maître d\'ouvrage : {META["mo"]}', small))
    story.append(Paragraph(f'Destinataire : {META["destinataire"]} &nbsp;|&nbsp; Réf. {META["ref"]} &nbsp;|&nbsp; Date : {META["date"]} &nbsp;|&nbsp; {META["devise"]}', small))
    story.append(Spacer(1, 4))

    # bloc métré
    met = [[Paragraph("<b>Données de métré (synthèse)</b>", small), "", ""]]
    for lib, val, src in METRE:
        met.append([Paragraph(lib, small), Paragraph(f"<b>{val}</b>", small), Paragraph(src, small)])
    tmet = Table(met, colWidths=[60*mm, 50*mm, 73*mm])
    tmet.setStyle(TableStyle([
        ("SPAN", (0,0), (2,0)),
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#EFEFEF")),
        ("LINEBELOW", (0,0), (-1,-1), 0.25, colors.HexColor("#DDDDDD")),
        ("TOPPADDING", (0,0), (-1,-1), 1), ("BOTTOMPADDING", (0,0), (-1,-1), 1),
    ]))
    story.append(tmet)
    story.append(Spacer(1, 4))
    story.append(Paragraph("⚠ <b>Quantités présumées (QP)</b> — établies sur photo aérienne, estimatif architecte et scan NavVis ; "
                           "à vérifier contradictoirement lors d'une visite sur site. Décompte final sur métré réel. "
                           "Les colonnes Prix unitaire et Total sont à compléter par le couvreur. "
                           "<b>Surface de reprise ≈ 50 m²</b> — relevé MO : &lt; ½ du versant Nord (~68 m²) + ⅓ du versant Est (~61 m²) ; "
                           "versants répartis depuis 237 m² selon l'orientation des arêtes (scan). À confirmer sur site.", warn))
    story.append(Paragraph("Type : QF=forfaitaire · QP=présumée · PG=poste global · SR=somme réservée · PM=pour mémoire. "
                           "Mesurage réf. CCTB (Wallonie) §34.1 Couvertures / §34.2 Étanchéités, NBN B 06-001. "
                           "<b>Variantes V1–V3 et options O1–O2 : alternatives à comparer, non additionnées au total.</b>", small))
    story.append(Spacer(1, 6))

    cw = [11*mm, 96*mm, 11*mm, 13*mm, 14*mm, 19*mm, 22*mm]
    header = [Paragraph(f"<b>{c}</b>", small) for c in COLS]
    data = [header]
    chap_rows, spans = [], []
    ri = 1
    for title, rows in CHAPTERS:
        data.append([Paragraph(f"<b>{title}</b>", chap), "", "", "", "", "", ""])
        chap_rows.append(ri); spans.append(("SPAN", (0, ri), (-1, ri))); ri += 1
        for n, d, t, u, q in rows:
            data.append([Paragraph(n, body), Paragraph(d, body), Paragraph(t, body),
                         Paragraph(u, body), Paragraph(qstr(q), body), "", ""]); ri += 1
    # récap
    for lab in ["TOTAL HTVA (ch. 0 à 7)", "TVA (6% / 21%)", "TOTAL TVAC"]:
        data.append(["", Paragraph(f"<b>{lab}</b>", body), "", "", "", "", ""])
        spans.append(("SPAN", (1, ri), (5, ri))); ri += 1

    t = Table(data, colWidths=cw, repeatRows=1)
    ts = [
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#1F3864")),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("GRID", (0,0), (-1,-1), 0.25, colors.HexColor("#CCCCCC")),
        ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("ALIGN", (2,0), (4,-1), "CENTER"),
        ("ALIGN", (5,0), (6,-1), "RIGHT"),
        ("TOPPADDING", (0,0), (-1,-1), 2), ("BOTTOMPADDING", (0,0), (-1,-1), 2),
        ("LEFTPADDING", (0,0), (-1,-1), 3), ("RIGHTPADDING", (0,0), (-1,-1), 3),
    ]
    for cr in chap_rows:
        ts.append(("BACKGROUND", (0,cr), (-1,cr), colors.HexColor("#1F3864")))
    for s in spans:
        ts.append(s)
    # récap rows shading
    for k in range(3):
        rr = len(data)-3+k
        ts.append(("BACKGROUND", (0,rr), (-1,rr), colors.HexColor("#FCE4D6")))
    t.setStyle(TableStyle(ts))
    story.append(t)

    # --- Variantes & options : NON additionnées au total ci-dessus ---
    story.append(Spacer(1, 6))
    vdata = [[Paragraph("<b>VARIANTES &amp; OPTIONS — à comparer, NON incluses dans le total ci-dessus</b>", chap),
              "", "", "", "", "", ""]]
    for n, d, t_, u, q in VARIANTES:
        vdata.append([Paragraph(n, body), Paragraph(d, body), Paragraph(t_, body),
                      Paragraph(u, body), Paragraph(qstr(q), body), "", ""])
    vt = Table(vdata, colWidths=cw)
    vt.setStyle(TableStyle([
        ("SPAN", (0,0), (-1,0)),
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#7F7F7F")),
        ("GRID", (0,0), (-1,-1), 0.25, colors.HexColor("#CCCCCC")),
        ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("ALIGN", (2,0), (4,-1), "CENTER"), ("ALIGN", (5,0), (6,-1), "RIGHT"),
        ("TOPPADDING", (0,0), (-1,-1), 2), ("BOTTOMPADDING", (0,0), (-1,-1), 2),
        ("LEFTPADDING", (0,0), (-1,-1), 3), ("RIGHTPADDING", (0,0), (-1,-1), 3),
    ]))
    story.append(vt)

    story.append(Spacer(1, 8))
    story.append(Paragraph(
        "Conditions à préciser par le soumissionnaire : délai d'exécution, validité de l'offre, "
        "garanties (décennale, produits), références, planning. Offre établie hors TVA ; "
        "TVA applicable selon le régime du maître d'ouvrage.", small))
    story.append(Spacer(1, 14))
    sig = Table([[Paragraph("Établi par (maître d'ouvrage) :", small), Paragraph("Lu et approuvé — l'entrepreneur (date, signature) :", small)]],
                colWidths=[88*mm, 88*mm])
    sig.setStyle(TableStyle([("LINEBELOW",(0,0),(-1,0),0.5,colors.white),
                             ("TOPPADDING",(0,0),(-1,0),18)]))
    story.append(sig)

    doc.build(story)
    return path

if __name__ == "__main__":
    out_csv  = build_csv(os.path.join(HERE, BASENAME + ".csv"))
    out_xlsx = build_xlsx(os.path.join(HERE, BASENAME + ".xlsx"))
    out_pdf  = build_pdf(os.path.join(HERE, BASENAME + ".pdf"))
    n = sum(len(rows) for _, rows in CHAPTERS)
    print(f"OK — {n} postes / {len(CHAPTERS)} chapitres + {len(VARIANTES)} variantes·options")
    for p in (out_csv, out_xlsx, out_pdf):
        print("  ->", os.path.relpath(p, os.path.dirname(HERE)), f"({os.path.getsize(p)} o)")
