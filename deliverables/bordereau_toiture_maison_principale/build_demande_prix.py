#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
DEMANDE DE PRIX — version simplifiée & ouverte, à ENVOYER aux toituriers.
Pure liste des travaux souhaités + colonnes « Votre proposition » et « Prix ».
Les MESURES/QUANTITÉS ne sont PLUS ici : elles sont dans la fiche annexe
« MESURES_releve_toiture » (build_mesures.py), jointe au même envoi.

Sort : DEMANDE_DE_PRIX_toiture_logis.pdf / .xlsx / .docx (éditable)
Régénérer : ./.venv/bin/python deliverables/bordereau_toiture_maison_principale/build_demande_prix.py
"""
import os
HERE = os.path.dirname(os.path.abspath(__file__))
BASE = "DEMANDE_DE_PRIX_toiture_logis"

TITRE = "Demande de prix — Toiture du logis (Ferme du Temple)"
SOUS  = "Ferme du Temple, Avenue Joseph Wauters 227, 7080 Frameries — collectif Habitat partagé"

INTRO = ("Bonjour, voici les travaux de toiture que nous aimerions faire chiffrer sur le logis "
         "(le bâtiment avec la mérule). La colonne de droite est là pour vos propositions : "
         "méthode, matériaux conseillés, alternatives, remarques — on est vraiment preneurs de "
         "vos conseils. Nos relevés (dimensions, surfaces et comment on les a obtenues) sont dans "
         "la fiche « Mesures » jointe ; ce sont des estimations, à confirmer ensemble sur place.")

# (titre de section, [ travaux en langage courant ])  — sans chiffres : voir fiche Mesures
SECTIONS = [
 ("Installation, échafaudage et accès", [
   "Échafaudage et accès en toiture (selon ce que vous jugez nécessaire)",
   "Protections et mise hors d'eau provisoire pendant le chantier",
 ]),
 ("Réparation de la couverture en ardoise (surtout les pans Nord et Est)", [
   "Reprise des zones d'ardoises abîmées / bâchées : dépose, sous-toiture et nouvelles ardoises. "
   "L'existant est de l'ardoise reconstituée — merci d'indiquer aussi, si possible, le prix en ardoise naturelle.",
   "Remplacement d'ardoises cassées isolées, ailleurs sur le toit",
   "Faîtage, arêtiers, noues et raccords (selon besoin)",
 ]),
 ("Petit toit plat (terrasson) en zinc, au sommet", [
   "Réfection de l'étanchéité du terrasson. Aujourd'hui en zinc ; on referait en zinc, mais votre "
   "avis nous intéresse (zinc / bitume / EPDM ?).",
 ]),
 ("Charpente", [
   "Vérification et réparations ponctuelles si nécessaire (chevrons, sablières) — à voir une fois "
   "la couverture ouverte. NB : un traitement contre la mérule est par ailleurs prévu, une "
   "coordination sera utile.",
 ]),
 ("Lucarnes, tourelle et cheminées", [
   "Lucarnes (toits cintrés en zinc) et œil-de-bœuf : reprise de la couverture et de la zinguerie",
   "Tourelle (toit conique) : révision",
   "Cheminées : solins et étanchéité au pourtour",
 ]),
 ("Évacuation des eaux de pluie", [
   "Gouttières (corniche) et descentes : révision / remplacement",
 ]),
 ("Divers et fin de chantier", [
   "Repérage amiante de la couverture avant dépose (obligatoire), évacuation des déchets, nettoyage",
 ]),
]

OPTIONS = [
 "Et si on refaisait toute la couverture à neuf ? (utile pour notre dossier de financement) — "
 "prix indicatif pour l'ensemble des pans, en ardoise reconstituée, et en naturelle si possible.",
]

PIED = ("Merci d'indiquer votre délai, la validité de l'offre et les garanties. Prix hors TVA. "
        "Un scan 3D du bâtiment est disponible (lien et accès dans notre message) : l'outil de "
        "mesure permet de prendre vos propres cotes. Voir aussi la fiche « Mesures » jointe.")

COLS = ["Travaux souhaités", "Votre proposition (méthode, matériaux, remarque)", "Prix € HTVA"]

# ============================ PDF ============================
def build_pdf(path):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

    st = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=st["Normal"], fontSize=9.5, leading=12)
    intro = ParagraphStyle("i", parent=st["Normal"], fontSize=9.5, leading=13)
    h1 = ParagraphStyle("h1", parent=st["Normal"], fontSize=14, leading=16, spaceAfter=2)
    sec = ParagraphStyle("sec", parent=st["Normal"], fontSize=10.5, leading=13, textColor=colors.HexColor("#333333"))
    small = ParagraphStyle("s", parent=st["Normal"], fontSize=8, leading=10, textColor=colors.HexColor("#666666"))

    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=16*mm, rightMargin=14*mm,
                            topMargin=14*mm, bottomMargin=14*mm, title=TITRE)
    s = [Paragraph(TITRE, h1), Paragraph(SOUS, small), Spacer(1, 6), Paragraph(INTRO, intro), Spacer(1, 8)]

    cw = [88*mm, 70*mm, 22*mm]
    data = [[Paragraph(f"<b>{c}</b>", small) for c in COLS]]
    extra = []
    ri = 1
    def section(title, fill):
        nonlocal ri
        data.append([Paragraph(f"<b>{title}</b>", sec), "", ""])
        extra.append(("SPAN", (0, ri), (-1, ri))); extra.append(("BACKGROUND", (0, ri), (-1, ri), colors.HexColor(fill)))
        ri += 1
    for title, items in SECTIONS:
        section(title, "#EFEFEF")
        for trav in items:
            data.append([Paragraph(trav, body), "", ""]); ri += 1
    section("Une question ouverte (si ça vous parle)", "#FDF3E7")
    for trav in OPTIONS:
        data.append([Paragraph(trav, body), "", ""]); ri += 1

    t = Table(data, colWidths=cw, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#5B7A9D")),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("GRID", (0,0), (-1,-1), 0.25, colors.HexColor("#CCCCCC")),
        ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#FAFAFA")]),
        ("TOPPADDING", (0,0), (-1,-1), 5), ("BOTTOMPADDING", (0,0), (-1,-1), 7),
        ("LEFTPADDING", (0,0), (-1,-1), 4), ("RIGHTPADDING", (0,0), (-1,-1), 4),
    ] + extra))
    s += [t, Spacer(1, 10), Paragraph(PIED, small)]
    doc.build(s); return path

# ============================ XLSX ============================
def build_xlsx(path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    wb = Workbook(); ws = wb.active; ws.title = "Demande de prix"
    thin = Side(style="thin", color="CCCCCC"); bd = Border(thin,thin,thin,thin)
    wrap = Alignment(wrap_text=True, vertical="top")
    for i,w in enumerate([64, 52, 16],1): ws.column_dimensions[get_column_letter(i)].width = w
    r = 1
    for text, fnt, hgt in [(TITRE, Font(bold=True, size=13), None),
                           (SOUS, Font(size=9, color="666666"), None),
                           (INTRO, Font(size=9, italic=True), 70)]:
        ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=3)
        c = ws.cell(r,1,text); c.font = fnt; c.alignment = wrap
        if hgt: ws.row_dimensions[r].height = hgt
        r += 1
    r += 1
    for i,name in enumerate(COLS,1):
        c = ws.cell(r,i,name); c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="5B7A9D")
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center"); c.border = bd
    r += 1
    def section(title, fill):
        nonlocal r
        ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=3)
        c = ws.cell(r,1,title); c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor=fill)
        for col in range(1,4): ws.cell(r,col).border = bd
        r += 1
    def item(trav):
        nonlocal r
        ws.cell(r,1,trav).alignment = wrap
        ws.cell(r,2,None).alignment = wrap; ws.cell(r,3,None)
        for col in range(1,4): ws.cell(r,col).border = bd
        ws.row_dimensions[r].height = 42; r += 1
    for title, items in SECTIONS:
        section(title, "EFEFEF")
        for trav in items: item(trav)
    section("Une question ouverte (si ça vous parle)", "FDF3E7")
    for trav in OPTIONS: item(trav)
    r += 1
    ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=3)
    c = ws.cell(r,1,PIED); c.font = Font(size=9, italic=True, color="666666"); c.alignment = wrap
    ws.row_dimensions[r].height = 55
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(path); return path

# ============================ DOCX (éditable) ============================
def build_docx(path):
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    def shade(cell, hexcolor):
        tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd')
        shd.set(qn('w:val'),'clear'); shd.set(qn('w:fill'),hexcolor); tcPr.append(shd)
    doc = Document(); sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE; sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for m in ("left_margin","right_margin","top_margin","bottom_margin"): setattr(sec, m, Cm(1.5))
    doc.add_heading(TITRE, level=1)
    doc.add_paragraph(SOUS).runs[0].italic = True
    doc.add_paragraph(INTRO); doc.add_paragraph()
    widths = [Cm(11.5), Cm(9.5), Cm(3.5)]
    table = doc.add_table(rows=1, cols=3); table.style = "Table Grid"; table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i,name in enumerate(COLS):
        cell = table.rows[0].cells[i]; cell.text = name; shade(cell, "5B7A9D")
        for run in cell.paragraphs[0].runs: run.font.bold = True; run.font.color.rgb = RGBColor(0xFF,0xFF,0xFF)
    def section(title, fill):
        row = table.add_row().cells
        a = row[0].merge(row[1]).merge(row[2]); a.text = title; shade(a, fill)
        a.paragraphs[0].runs[0].font.bold = True
    def item(trav):
        table.add_row().cells[0].text = trav
    for title, items in SECTIONS:
        section(title, "EFEFEF")
        for trav in items: item(trav)
    section("Une question ouverte (si ça vous parle)", "FDF3E7")
    for trav in OPTIONS: item(trav)
    for row in table.rows:
        for i,c in enumerate(row.cells): c.width = widths[i]
    doc.add_paragraph()
    p = doc.add_paragraph(PIED); p.runs[0].italic = True
    for run in p.runs: run.font.size = Pt(9)
    doc.save(path); return path

if __name__ == "__main__":
    outs = [build_pdf(os.path.join(HERE, BASE+".pdf")),
            build_xlsx(os.path.join(HERE, BASE+".xlsx")),
            build_docx(os.path.join(HERE, BASE+".docx"))]
    n = sum(len(it) for _,it in SECTIONS)
    print(f"OK — demande de prix (sans chiffres) : {n} lignes de travaux + {len(OPTIONS)} question(s) ouverte(s)")
    for p in outs: print("  ->", os.path.basename(p), f"({os.path.getsize(p)} o)")
