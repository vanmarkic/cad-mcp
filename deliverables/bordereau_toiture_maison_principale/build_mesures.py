#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
FICHE MESURES — annexe à joindre à la demande de prix.
Regroupe nos relevés (dimensions, surfaces) + en clair comment on les a obtenus.
Langage simple (pas de CCTB/QF-QP : ça reste dans METHODOLOGIE_metres.md, interne).

Sort : MESURES_releve_toiture.pdf / .docx (éditable)
Régénérer : ./.venv/bin/python deliverables/bordereau_toiture_maison_principale/build_mesures.py
"""
import os
HERE = os.path.dirname(os.path.abspath(__file__))
BASE = "MESURES_releve_toiture"

TITRE = "Mesures — Toiture du logis (Ferme du Temple)"
SOUS  = "Annexe à la demande de prix — collectif Habitat partagé de la Ferme du Temple"
INTRO = ("Voici nos relevés et comment on les a obtenus. Ce sont des estimations (scan 3D + "
         "estimatif de l'architecte + nos mesures) : des ordres de grandeur, à confirmer "
         "ensemble lors d'une visite. Le décompte final se ferait sur les mesures réelles.")

# (mesure, valeur, comment obtenue)
MESURES = [
 ("Surface du toit (versants en ardoise)", "≈ 237 m²", "estimatif de l'architecte, recoupé avec le scan 3D"),
 ("Emprise au sol du bâtiment", "182 m²", "scan 3D (NavVis / immovision)"),
 ("Pourtour du bâtiment", "≈ 58 m", "scan 3D"),
 ("  · dont corniche libre (gouttières)", "≈ 34 m", "côtés non mitoyens (les autres ~25 m sont contre l'aile ouest et l'atelier)"),
 ("Hauteur de façade jusqu'à la corniche", "≈ 10 m", "relevé sur le scan 3D (utile pour l'échafaudage)"),
 ("Zones d'ardoises à reprendre", "≈ 50 m²", "notre relevé : moins de la moitié du pan Nord + un tiers du pan Est"),
 ("Petit toit plat (terrasson zinc)", "≈ 20 m² (5 × 4 m)", "mesure à la main sur le scan 3D (outil de mesure immovision)"),
 ("Éléments en toiture", "1 tourelle · 2 lucarnes + 1 œil-de-bœuf · ~3 cheminées", "comptés sur les photos drone"),
 ("Couverture actuelle", "ardoise reconstituée", "—"),
 ("Pente du toit", "≈ 40° (à confirmer)", "déduite de la surface et de l'emprise"),
]

COMMENT = [
 "Le bâtiment a été scanné en 3D (NavVis / immovision) : l'emprise au sol, le pourtour et les "
 "hauteurs en sortent directement (mesures fiables).",
 "Les surfaces de toiture viennent de l'estimatif de l'architecte, recoupées avec le scan "
 "(cohérentes pour une pente d'environ 40°).",
 "Les zones à reprendre correspondent à notre observation des parties bâchées / abîmées "
 "(surtout les pans Nord et Est). À affiner une fois le toit ouvert.",
 "Le terrasson (petit toit plat) a été mesuré à la main sur le scan 3D : environ 5 × 4 m.",
 "Le nombre de lucarnes, la tourelle et les cheminées sont comptés sur les photos drone.",
 "Tout est indicatif et à confirmer sur place (et avec le géomètre, dont les relevés précis suivront).",
]
PIED = ("Scan 3D accessible en ligne (lien et accès dans notre message) — l'outil « mesure / cotes "
        "à la volée » permet de prendre vos propres cotes directement sur le modèle.")

# ============================ PDF ============================
def build_pdf(path):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, ListFlowable, ListItem
    st = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=st["Normal"], fontSize=9.5, leading=12)
    h1 = ParagraphStyle("h1", parent=st["Normal"], fontSize=14, leading=16, spaceAfter=2)
    h2 = ParagraphStyle("h2", parent=st["Normal"], fontSize=11, leading=13, spaceBefore=8, spaceAfter=4)
    small = ParagraphStyle("s", parent=st["Normal"], fontSize=8, leading=10, textColor=colors.HexColor("#666666"))
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=18*mm, rightMargin=16*mm, topMargin=14*mm, bottomMargin=14*mm, title=TITRE)
    s = [Paragraph(TITRE, h1), Paragraph(SOUS, small), Spacer(1,6), Paragraph(INTRO, body), Spacer(1,8)]
    data = [[Paragraph("<b>Mesure</b>", small), Paragraph("<b>Valeur</b>", small), Paragraph("<b>Comment on l'a obtenue</b>", small)]]
    for m,v,c in MESURES:
        data.append([Paragraph(m, body), Paragraph(f"<b>{v}</b>", body), Paragraph(c, small)])
    t = Table(data, colWidths=[62*mm, 42*mm, 76*mm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#5B7A9D")), ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("GRID", (0,0), (-1,-1), 0.25, colors.HexColor("#CCCCCC")), ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#FAFAFA")]),
        ("TOPPADDING", (0,0), (-1,-1), 3), ("BOTTOMPADDING", (0,0), (-1,-1), 4),
    ]))
    s += [t, Spacer(1,6), Paragraph("Comment on a mesuré", h2)]
    s.append(ListFlowable([ListItem(Paragraph(x, body), leftIndent=10) for x in COMMENT], bulletType="bullet", start="•"))
    s += [Spacer(1,8), Paragraph(PIED, small)]
    doc.build(s); return path

# ============================ DOCX (éditable) ============================
def build_docx(path):
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    def shade(cell, hexcolor):
        tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd')
        shd.set(qn('w:val'),'clear'); shd.set(qn('w:fill'),hexcolor); tcPr.append(shd)
    doc = Document()
    for m in ("left_margin","right_margin","top_margin","bottom_margin"): setattr(doc.sections[0], m, Cm(1.8))
    doc.add_heading(TITRE, level=1)
    doc.add_paragraph(SOUS).runs[0].italic = True
    doc.add_paragraph(INTRO); doc.add_paragraph()
    table = doc.add_table(rows=1, cols=3); table.style = "Table Grid"
    for i,name in enumerate(["Mesure","Valeur","Comment on l'a obtenue"]):
        cell = table.rows[0].cells[i]; cell.text = name; shade(cell, "5B7A9D")
        for run in cell.paragraphs[0].runs: run.font.bold = True; run.font.color.rgb = RGBColor(0xFF,0xFF,0xFF)
    for m,v,c in MESURES:
        row = table.add_row().cells
        row[0].text = m; row[1].text = v; row[2].text = c
        for run in row[1].paragraphs[0].runs: run.font.bold = True
    widths = [Cm(6.5), Cm(4.5), Cm(7.5)]
    for row in table.rows:
        for i,c in enumerate(row.cells): c.width = widths[i]
    doc.add_paragraph()
    doc.add_heading("Comment on a mesuré", level=2)
    for x in COMMENT: doc.add_paragraph(x, style="List Bullet")
    doc.add_paragraph()
    p = doc.add_paragraph(PIED); p.runs[0].italic = True
    for run in p.runs: run.font.size = Pt(9)
    doc.save(path); return path

if __name__ == "__main__":
    outs = [build_pdf(os.path.join(HERE, BASE+".pdf")), build_docx(os.path.join(HERE, BASE+".docx"))]
    print(f"OK — fiche Mesures : {len(MESURES)} mesures + {len(COMMENT)} points de méthode")
    for p in outs: print("  ->", os.path.basename(p), f"({os.path.getsize(p)} o)")
